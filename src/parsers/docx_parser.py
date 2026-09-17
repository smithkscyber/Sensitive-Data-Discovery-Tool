"""DOCX extraction: paragraphs, then any table cells."""

from __future__ import annotations

from pathlib import Path

import docx


def extract(path: Path | str) -> str:
    """Return a Word document's text, one paragraph per line.

    Table content is walked separately because ``Document.paragraphs`` does not
    include it -- a detector that skipped tables would miss exactly the kind of
    dense, structured record that holds the most PII.
    """
    document = docx.Document(str(path))

    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(", ".join(cell.text for cell in row.cells))

    return "\n".join(lines)
